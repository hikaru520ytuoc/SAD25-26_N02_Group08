from pathlib import Path
import re,subprocess,sys
from docx import Document
from docx.shared import Inches,Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
r=Path(__file__).resolve().parents[1];logs=r/'results';out=r/'docs/BaoCao_ChuDe5_ThucNghiem.docx'
blue='19395B';pale='EAF1F8'
D=Document();sec=D.sections[0]
sec.top_margin=Cm(2.0);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.2)
normal=D.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5);normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.25
for hn,size in [('Title',19),('Heading 1',15),('Heading 2',12.5)]:
 st=D.styles[hn];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(blue)
 st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(6)
code=D.styles.add_style('SourceCode',WD_STYLE_TYPE.PARAGRAPH);code.font.name='Consolas';code.font.size=Pt(8.6);code.paragraph_format.line_spacing=1.0;code.paragraph_format.space_before=Pt(3);code.paragraph_format.space_after=Pt(5)

# footer
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Chủ đề 5 — Software Security  |  Trang ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(t='',bold=False,ital=False):
 q=D.add_paragraph();rr=q.add_run(t);rr.bold=bold;rr.italic=ital;return q

def h(t,n=1):return D.add_heading(t,level=n)

def tbl(headers,rows,widths=None):
 t=D.add_table(rows=1, cols=len(headers));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for i,a in enumerate(headers):
  cell=t.rows[0].cells[i];cell.text=str(a)
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),blue);cell._tc.get_or_add_tcPr().append(shade)
  for run in cell.paragraphs[0].runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255);run.font.size=Pt(9)
 for j,data in enumerate(rows):
  cells=t.add_row().cells
  for i,a in enumerate(data):
   cells[i].text=str(a)
   for para in cells[i].paragraphs:
    para.paragraph_format.space_after=Pt(1)
    for run in para.runs:run.font.size=Pt(8.7)
   if j%2==1:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),pale);cells[i]._tc.get_or_add_tcPr().append(shade)
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths):row.cells[i].width=Cm(w)
 D.add_paragraph('')
 return t

def pre(text,caption=None,limit=32):
 if caption:p(caption,bold=True)
 for ix,line in enumerate(text.strip('\n').splitlines()[:limit]):
  q=D.add_paragraph(style='SourceCode');q.paragraph_format.space_after=Pt(0)
  q.add_run(line.replace('\t','    '))
 D.add_paragraph('')

def page():D.add_page_break()

# Cover
for t in ['ĐẠI HỌC PHENIKAA','HỌC PHẦN CSE703093 — AN TOÀN PHẦN MỀM']:
 q=p(t);q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].bold=True
D.add_paragraph('\n')
q=D.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;r0=q.add_run('BÁO CÁO BÀI TẬP LỚN');r0.bold=True;r0.font.size=Pt(20);r0.font.color.rgb=RGBColor.from_string(blue)
q=p('CHỦ ĐỀ 5: QUY TRÌNH END-TO-END TÍCH HỢP 4 CHƯƠNG');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].bold=True
q=p('Kiểm chứng an toàn chương trình nhập điểm đồ án tốt nghiệp bằng CSV');q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.runs[0].font.size=Pt(14)
D.add_paragraph('\n')
tbl(['Thông tin','Nội dung'],[
 ['Chương trình mục tiêu','thesis_grade_parser.c — C11'],
 ['Nhóm','Nhóm 06 (bổ sung MSSV nếu cần)'],
 ['Thành viên','Nguyễn Cao Hải; Ngô Nhật Quang; Đồng Đại Đạt'],
 ['Giảng viên','ThS. Vũ Quang Dũng'],
 ['Phạm vi','Mô hình hóa → Phân tích tĩnh/sửa bộ nhớ → Z3/BMC → Fuzzing'],
 ['Tình trạng','Đã chạy kiểm thử và lưu log; chưa chạy ứng dụng Windows Spec Verification Lab'],
])
q=p('Hà Nội, tháng 10 năm 2026');q.alignment=WD_ALIGN_PARAGRAPH.CENTER
page()

h('MỤC LỤC')
for t in ['1. Giới thiệu và mục tiêu','2. Chương trình mục tiêu và đặc tả','3. Phương pháp và tổ chức pipeline (Chương 1–4)','4. Kết quả kiểm chứng và bản vá','5. Đánh giá độ bao phủ và thử nghiệm fuzzing','6. Thảo luận, giới hạn và kết luận','7. Phụ lục: hướng dẫn chạy lại, log và tài liệu tham khảo']:
 p(t)
p('Lưu ý: các tiêu đề “Chương 1–4” ở đây tương ứng bốn kỹ thuật của học phần; không tạo “Chương 5 – Phân loại CWE/CVE” theo mẫu sinh tự động cũ.',ital=True)
h('TÓM TẮT')
p('Báo cáo kiểm tra một chương trình C nhỏ đọc dữ liệu CSV nhập điểm đồ án. Hai phiên bản của cùng một chương trình là v0 (cài lỗi có chủ đích trong sandbox) và v1 (đã sửa). Quá trình sử dụng mô hình Kripke với kiểm tra đường đi hữu hạn, công cụ quét mẫu của softsec-toolkit, Clang AddressSanitizer/UndefinedBehaviorSanitizer/LeakSanitizer, thư viện Z3 và kiểm thử ngẫu nhiên có seed cố định. Các kết quả đã chạy được liên kết tới file log, thay vì suy đoán hay tạo log mô phỏng.')
p('Từ khóa: CSV parser; Kripke; BMC; CWE-120; CWE-401; CWE-787; SMT; fuzzing.')
page()

h('1. Giới thiệu và mục tiêu')
p('Trong hệ thống Quản lý đồ án tốt nghiệp, dữ liệu điểm thường được nhập hoặc tổng hợp từ nhiều giảng viên. Một chức năng nhập CSV sai kiểm tra kích thước có thể gây ghi vượt vùng nhớ, truy cập mảng sai vị trí hoặc rò rỉ vùng nhớ khi dữ liệu không hợp lệ. Đề tài tập trung vào trust boundary là file CSV do người dùng cung cấp, không phân tích toàn bộ web NestJS/Next.js.')
p('Mục tiêu: (i) đặc tả luồng và điều kiện an toàn; (ii) phát hiện, tái hiện và sửa ít nhất hai lỗi bộ nhớ; (iii) kiểm chứng một thuộc tính chỉ số mảng bằng Z3; (iv) chạy fuzzing xác nhận sau khi sửa; (v) ghi log và trình bày các hạn chế của kiểm chứng.')
p('Phạm vi: chương trình C11 xử lý 6 trường CSV, 100 bản ghi tối đa; không bao gồm cơ sở dữ liệu, xác thực giảng viên, tính trọng số điểm hay nhập CSV có dấu phẩy nằm trong chuỗi trích dẫn.')
h('2. Chương trình mục tiêu và đặc tả')
p('Đầu vào là file CSV, đầu ra là “ACCEPT: N records” với mã thoát 0 hoặc thông báo “REJECT” với mã thoát 2. Điểm được nhập ở ba cột nguồn, không nhập điểm tổng kết.')
tbl(['Cột CSV','Giới hạn','Ý nghĩa'],[
 ['student_id','tối đa 15 byte','Mã sinh viên'],
 ['student_name','tối đa 63 byte','Tên sinh viên'],
 ['project_id','tối đa 15 byte','Mã đồ án'],
 ['supervisor_score','0–10, hữu hạn','Điểm hướng dẫn'],
 ['reviewer_score','0–10, hữu hạn','Điểm phản biện'],
 ['council_score','0–10, hữu hạn','Điểm hội đồng'],
])
pre('student_id,student_name,project_id,supervisor_score,reviewer_score,council_score\nSV001,Nguyen Van A,DA001,8.5,8.0,9.0', 'Listing 1. Ví dụ CSV đầu vào hợp lệ')
p('Luồng xử lý: đọc header → tách đủ sáu trường → sao chép vào GradeRecord → kiểm tra điểm → lưu records[record_count] nếu còn dung lượng → chấp nhận/từ chối. Giới hạn MAX_RECORDS = 100, STUDENT_NAME_BYTES = 64, MAX_LINE_BYTES = 512.')
tbl(['ID','Điều kiện kiểm chứng','Công cụ'],[
 ['REQ-03','Mọi phép ghi mảng có index < 100','BMC mô hình + Z3 + ASan'],
 ['REQ-05','Không sao chép quá giới hạn mảng char','Static scanner + ASan'],
 ['REQ-06','Không rò rỉ malloc trên nhánh lỗi','LSan'],
 ['REQ-07','Tránh tràn uint32 khi tính kích thước','Z3 + chạy --estimate'],
 ['REQ-08','Input sai phải từ chối an toàn','Regression + fuzzing'],
])
page()

h('3. Phương pháp và tổ chức pipeline')
p('Một chương trình duy nhất được duy trì dưới hai bản nguồn: src/thesis_grade_parser_v0.c (có chủ đích cài lỗi) và src/thesis_grade_parser_v1.c (đã vá). Các script tự sinh hai bản từ cùng source_original.c để dễ truy vết. File JSON config/spec.json sinh ra src/parser_config.h; Excel đặc tả đi kèm dùng nhập vào phần mềm của thầy, không được nhận là đã chạy qua ứng dụng Windows.')
h('3.1. Chương 1 — Mô hình hóa Kripke',2)
p('Dùng lớp KripkeStructure lấy từ softsec-toolkit/ch1_models/kripke.py. Mô hình v0 có thêm đường đi parsing → unsafe_store không qua validated; v1 loại bỏ đường đi đó. Các trạng thái cuối success, error tự lặp để tránh deadlock giả do không khai báo quan hệ đi ra.')
tbl(['Trạng thái','Nhãn chính','Dịch chuyển tiêu biểu'],[
 ['init','init','read_header'],['read_header','header','parsing / error'],['parsing','parsing','validated / error / unsafe_store (chỉ v0)'],['validated','validated','stored'],['stored','stored, validated','parsing / success'],['unsafe_store','violation, stored','error (chỉ v0)'],['success; error','success; error','tự lặp'],
])
p('Thuộc tính cần kiểm tra: AG(stored → validated), AG(error → ¬stored), EF(success). Việc tìm đường đi tới trạng thái lưu chưa xác thực được thực hiện bằng duyệt đồ thị hữu hạn theo giới hạn k, không phải tự động suy diễn toàn bộ semantics của chương trình C.')
p('Mã TikZ (LaTeX) — dán trực tiếp vào tài liệu LaTeX có \\usepackage{tikz} để biên dịch thành hình vector, thay cho ảnh tĩnh:')
pre(r'''\begin{tikzpicture}[node distance=15mm,->,>=stealth]
\node (a) {init};
\node[right of=a] (b) {read\_header};
\node[right of=b] (c) {parsing};
\node[below of=c] (d) {validated};
\node[left of=d] (e) {stored};
\node[below of=e] (f) {success};
\draw (a)--(b); \draw (b)--(c);
\draw (c)--(d); \draw (d)--(e); \draw (e)--(f);
\end{tikzpicture}''')
h('3.2. Chương 2 — Static analysis và sửa lỗi bộ nhớ',2)
p('Mã C sinh từ đặc tả',bold=True)
p('Các khuôn mẫu sinh mã từ workbook được dùng minh họa; module xử lý CSV thực tế được viết theo đặc tả và tách v0/v1 để đối chiếu. Không coi mã C do công cụ mẫu sinh ra là toàn bộ chương trình.',ital=True)
p('Sử dụng chính pattern_scanner.py từ softsec-toolkit/ch2_memsafe; sau đó biên dịch Clang với -fsanitize=address,undefined và chạy file long_name.csv, records_101.csv, invalid_score.csv. Quét mẫu chỉ chỉ ra lời gọi strcpy có rủi ro; sanitizer cung cấp bằng chứng động về hành vi khi chạy.')
pre('clang -std=c11 -O0 -g -fsanitize=address,undefined \\\n  src/thesis_grade_parser_v0.c -o build/parser_v0', 'Listing 2. Lệnh biên dịch kiểm thử v0')
h('3.3. Chương 3 — Kiểm chứng SMT/BMC',2)
p('Thuộc tính chính: 0 ≤ record_count < 100 trước khi ghi records[record_count]. Z3 nhận mô hình hóa số nguyên của điều kiện nhánh đã trích thủ công từ store_record(); phiên bản v0 bỏ guard còn phiên bản v1 kiểm tra guard. Đây là SMT cho biểu thức của mã nguồn chứ không phải chứng minh tổng thể mọi hoạt động của C.')
pre('(declare-const idx Int)\n(assert (>= idx 0))\n(assert (< idx 100))\n(assert (>= idx 100))\n(check-sat)', 'Listing 3. Điều kiện v1 yêu cầu vừa index hợp lệ vừa index vượt biên → UNSAT')
p('Ở phần tràn số, kiểm tra count × sizeof(GradeRecord) vượt UINT32_MAX. Kích thước bản ghi trên môi trường thử nghiệm là 108 byte; sau khi thay đổi cấu trúc cần xác nhận lại sizeof thực tế.')
h('3.4. Chương 4 — Fuzzing và xác nhận',2)
p('Kiểm thử hộp đen tạo 300 file CSV đột biến bằng seed 20261011, kiểm tra return code và timeout. Kiểm thử định hướng mã (white-box branch-directed) giải điều kiện index >= 100 bằng Z3 và tạo đúng file 101 bản ghi; kiểm tra v1 từ chối. Đây không phải engine symbolic execution tự động như KLEE/angr.')
page()

h('4. Kết quả kiểm chứng và bản vá')
h('4.1. CWE-120 — Sao chép tên không kiểm tra chiều dài',2)
p('Vị trí: copy_field(), v0 dòng 58. Nguyên nhân: strcpy() không biết kích thước bộ đệm đích; tên dài 120 byte gây ghi tràn scratch->student_name[64].')
pre('/* v0 */\nstrcpy(dst, src);\n\n/* v1 */\nsize_t n = strlen(src);\nif (!n || n >= capacity) return 0;\nmemcpy(dst, src, n + 1);', 'Listing 4. Sửa lỗi trước/sau')
p('Bằng chứng v0: results/05_v0_long_name.log ghi “AddressSanitizer: heap-buffer-overflow”, WRITE of size 121. Bản v1 từ chối tại dòng thứ hai của CSV và trong chính ca kiểm tra này không có báo cáo ASan.')
h('4.2. CWE-787 — Ghi vượt mảng bản ghi',2)
p('Vị trí: store_record(), v0 dòng 87. Nguyên nhân: v0 thực hiện records[record_count++] dù record_count đã bằng 100, vượt mảng chứa đúng 100 phần tử. Lỗi được xác nhận bằng UBSan và ASan khi nhập file 101 bản ghi.')
pre('/* v0 */\nrecords[record_count++] = record;\n\n/* v1 */\nif (record_count >= MAX_RECORDS) return 0;\nrecords[record_count++] = record;', 'Listing 5. Điều kiện kiểm tra giới hạn trước khi ghi')
p('Bằng chứng v0: results/05_v0_records_101.log có “index 100 out of bounds” và “global-buffer-overflow”. Bản v1 chấp nhận đúng 100 bản ghi, từ chối 101 bản ghi (mã 2).')
h('4.3. CWE-401 — Rò rỉ bộ nhớ trên nhánh lỗi',2)
p('Vị trí: parse_record() trên nhánh điểm invalid. V0 trả về khi score = nan mà không free(scratch), gây leak. V1 giải phóng scratch trước khi trả về.')
pre('/* v0 */\nif (!parse_score(...)) return 0;\n\n/* v1 */\nif (!parse_score(...)) {\n    free(scratch);\n    return 0;\n}', 'Listing 6. Đảm bảo giải phóng trên nhánh thất bại')
p('Bằng chứng: results/05_v0_invalid_score.log có “LeakSanitizer: detected memory leaks”, Direct leak of 108 byte(s). Bản v1 từ chối cùng file CSV và không có leak report trong lần kiểm thử này.')
h('4.4. CWE-190 — Nguy cơ tràn số nguyên không dấu',2)
p('V0 tính uint32_t bytes = count * sizeof(GradeRecord) không guard. Dù phép nhân unsigned wraparound có định nghĩa trong C, giá trị byte bị sai và nguy hiểm nếu dùng để cấp phát thật. V1 kiểm tra count > UINT32_MAX / sizeof(GradeRecord) trước khi nhân. Đối số thử: --estimate 4294967295.')

h('4.5. Bảng tổng hợp before/after',2)
tbl(['Lỗi','v0 (minh chứng)','v1 (minh chứng)','Công cụ'],[
 ['CWE-120','ASan heap-buffer-overflow','REJECT long_name, không báo ASan ở ca này','pattern_scanner + ASan'],
 ['CWE-787','UBSan index 100; ASan global-buffer-overflow','ACCEPT 100 / REJECT 101','ASan/UBSan + Z3'],
 ['CWE-401','LSan leak 108 byte','REJECT nan, không báo LSan ở ca này','LeakSanitizer'],
 ['CWE-190','count × size bị wrap','Từ chối overflow count','--estimate + Z3'],
])
p('Không có báo cáo sanitizer trong một tập test hữu hạn không có nghĩa toàn bộ mã nguồn đã được chứng minh an toàn.')
h('4.6. Kết quả BMC/Z3 đã chạy',2)
pre((logs/'03_kripke_bmc.log').read_text(), 'Listing 7. Trích log Kripke/BMC (mô hình trừu tượng)',limit=15)
pre((logs/'06_smt_z3.log').read_text(), 'Listing 8. Log Z3 thật thông qua libz3.so (SMT-LIB2)',limit=12)
p('SAT ở v0 đưa ra index=100 làm nhân chứng. UNSAT ở v1 khẳng định tổ hợp điều kiện idx<100 và idx>=100 là không thỏa mãn trong mô hình nguyên này. Đối với BMC k≤10, không tìm thấy unsafe_store ở v1; điều này chỉ kết luận trên mô hình và giới hạn đã kiểm tra.')
h('5. Đánh giá độ bao phủ và fuzzing')
fuzz=(logs/'07_fuzz.log').read_text();print_fuzz=fuzz.splitlines()[0] if fuzz else ''
p(print_fuzz)
tbl(['Hạng mục','Giá trị quan sát'],[
 ['Số ca fuzzing hộp đen','300 ca, seed 20261011'],
 ['Đầu vào chấp nhận','137'],['Đầu vào bị từ chối','163'],
 ['Thoát bất thường / timeout','0 / 0'],
 ['GCOV số dòng thực thi','93,41% trên 91 dòng tính được'],
 ['GCOV hướng nhánh đi qua ít nhất một lần','73,08% trên 78 hướng nhánh'],
])
p('Đây là số liệu thực nghiệm trên môi trường hiện tại, không phải ước lượng từ tài liệu. Số lượng test không tương đương coverage. “Branches executed 97,44%” trong gcov khác với “Taken at least once 73,08%”: báo cáo sử dụng chỉ số thứ hai để diễn giải mức bao phủ hướng nhánh.')
pre((logs/'09_regression.log').read_text(),'Listing 9. Một phần 20 bài kiểm thử hồi quy',limit=14)
p('White-box branch-directed tạo 101 bản ghi theo witness idx=100 của Z3; chạy file đó trên v1 thu được REJECT: >100 records. Đây là kiểm thử định hướng đường đi có sử dụng điều kiện mã, không phải white-box concolic engine tự động.')
page()

h('6. Thảo luận, hạn chế và kết luận')
h('6.1. Điểm đã đạt',2)
p('Nhóm đã xây dựng một đối tượng thực nghiệm C gắn với nghiệp vụ quản lý đồ án, có nguồn v0/v1 và đặc tả Excel. Một chuỗi duy nhất thu thập bằng chứng từ mô hình hóa tới phân tích tĩnh, sanitizer, Z3 và fuzzing. Các lỗi tràn bộ đệm, vượt mảng, rò rỉ bộ nhớ có log trước/sau; thuộc tính index được kiểm tra ở hai phiên bản bằng SMT.')
h('6.2. Hạn chế còn lại',2)
p('Đầu vào CSV chỉ hỗ trợ 6 trường không có dấu phẩy trích dẫn; tên/ID được giới hạn theo byte, chưa xử lý Unicode normalize. Việc xác nhận v1 chỉ áp dụng cho các testcase đã chạy. Model Kripke và công thức SMT là phép trừu tượng hóa/tách điều kiện do nhóm viết, chưa phải semantic proof toàn chương trình. BMC chỉ xét k≤10 và không chứng minh tính an toàn vô hạn.')
p('Chương trình không kết nối backend thật. 300 phép thử với seed cố định có thể bỏ sót những nhánh hiếm. Coverage 73,08% hướng nhánh cho thấy vẫn còn nhánh chưa được kiểm thử. Bài chưa thực thi ứng dụng Windows Spec Verification Lab; file Excel mới được kiểm tra cấu trúc và tên cột. Các hàm sinh từ Excel trong ứng dụng của thầy có thể là hàm minh họa theo template, không phải tự động toàn bộ module C đã viết.')
h('6.3. Hướng hoàn thiện',2)
p('Bổ sung chạy trên máy nhóm với đúng môi trường thầy sử dụng, xuất log của Spec Verification Lab, kiểm tra các testcase chưa phủ, tăng seed và thời gian fuzzing, và nếu cần mới mở rộng import JSON tới NestJS. Không đưa bản v0 vào bất kỳ API hay hệ thống thật nào.')
h('6.4. Kết luận',2)
p('Với phạm vi một môn học, pipeline minh họa rõ lợi ích bổ sung lẫn nhau giữa mô hình hóa, phân tích tĩnh, kiểm chứng ràng buộc và kiểm thử động. Việc vá và chạy lại bằng chứng quan trọng hơn một kết luận “an toàn 100%” không có cơ sở.')
h('7. Phụ lục và hướng dẫn tái lập')
h('7.1. Cài đặt và chạy lại',2)
pre('cd ATPM_ChuDe5_ThucThi\nchmod +x scripts/run_pipeline.sh\n./scripts/run_pipeline.sh\n# Xem ket qua:\nls results/*.log\ncat results/06_smt_z3.log\ncat results/08_gcov.log','Listing 10. Lệnh chạy trên Ubuntu có clang, gcc, Python3 và libz3')
p('Điều kiện môi trường: GNU/Linux, Python 3, Clang hỗ trợ ASan/UBSan/LSan, GCC/gcov và thư viện chia sẻ libz3.so. Không cần cài z3-solver Python vì script sử dụng C API của Z3 qua ctypes.')
h('7.2. Thư mục bằng chứng',2)
tbl(['Đường dẫn','Ý nghĩa'],[
 ['results/03_kripke_bmc.log','Mô hình v0/v1, độ sâu k và counterexample'],
 ['results/04_static_scan.log','Phân tích mẫu C v0/v1'],
 ['results/05_v0_*.log và 05_v1_*.log','ASan/UBSan/LSan trước-sau'],
 ['results/06_smt_z3.log','SAT/UNSAT và witness trên libz3'],
 ['results/07_fuzz.log','Số ca, seed, kết quả mutation và boundary'],
 ['results/08_gcov.log','Line/branch coverage'],
 ['results/09_regression.log','Bộ ca hồi quy'],
])
h('7.3. Tài liệu sử dụng',2)
p('[1] ThS. Vũ Quang Dũng, BTL_01_HuongDan_DeTai.pdf, CSE703093 — An toàn phần mềm, tài liệu nhóm cung cấp.')
p('[2] ThS. Vũ Quang Dũng, BTL_02_HuongDan_BaoCao.pdf, CSE703093 — An toàn phần mềm, tài liệu nhóm cung cấp.')
p('[3] Spec Verification Lab, Hướng dẫn tạo Spec file dạng excel.pdf, tài liệu nhóm cung cấp.')
p('[4] softsec-toolkit.zip: ch1_models/kripke.py, ch2_memsafe/pattern_scanner.py và ví dụ ch3_verify/ch4_dynafuzz, tài liệu thực hành nhóm cung cấp.')
p('[5] Z3 C API, Clang sanitizers, GCC gcov — phần mềm được dùng trong thực nghiệm.')
p('Ghi chú liêm chính: bản thiết kế, báo cáo và code có sự hỗ trợ từ AI; nhóm cần kiểm tra, tái lập kết quả và ghi rõ mức độ hỗ trợ khi nộp chính thức.')

D.save(out)
print('WRITTEN',out,out.stat().st_size)
